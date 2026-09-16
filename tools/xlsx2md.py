# -*- coding: utf-8 -*-
import openpyxl, os, sys

def esc(v):
    if v is None: return ''
    return str(v).replace('\n', '<br>').replace('|', '\\|').strip()

def sheet_md(ws, title):
    rows = [[esc(c) for c in r] for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(x for x in r)]
    if not rows: return ''
    out = [f'## {title}', '']
    ncol = max(len(r) for r in rows)
    if ncol == 1:
        for r in rows:
            if r[0]: out.append(r[0]); out.append('')
        return '\n'.join(out)
    head = (rows[0] + [''] * ncol)[:ncol]
    out.append('| ' + ' | '.join(head) + ' |')
    out.append('|' + '|'.join([' --- '] * ncol) + '|')
    for r in rows[1:]:
        r = (r + [''] * ncol)[:ncol]
        out.append('| ' + ' | '.join(x if x else ' ' for x in r) + ' |')
    out.append('')
    return '\n'.join(out)

def convert(path, meta, title):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = ['---']
    for k, v in meta.items(): out.append(f'{k}: {v}')
    out += ['---', '', f'# {title}', '']
    for s in wb.sheetnames:
        md = sheet_md(wb[s], s)
        if md: out.append(md)
    return '\n'.join(out).rstrip() + '\n'

if __name__ == '__main__':
    BASE = os.path.expanduser('~/mnt/bi_datawarehouse')
    OUT = os.path.join(BASE, 'LLD')
    JOBS = [
     ('LD-05A_报表与仪表板规格册.xlsx','LD-05A-report-specs.md','LD-05A','报表与仪表板规格册','L5',
      '40 张报表与 20 个仪表板的规格册。现为样板与骨架，非定稿——须在合同签订后与客户逐张对齐'),
     ('HD-04_约束与设计决策登记册.xlsx','REF-HD-04-decisions.md','HD-04','约束与设计决策登记册','参考',
      '36 条约束与 61 项设计决策，含被拒方案与拒绝理由。改动任何架构决定之前先读这一份'),
     ('HD-02_需求追溯矩阵.xlsx','REF-HD-02-traceability.md','HD-02','需求追溯矩阵','参考',
      '189 条 RFP 条款 → 124 条需求 → HLD 章节 → LD 册章的逐条对应'),
    ]
    for src, dst, bid, title, layer, summary in JOBS:
        meta = {'id': bid, 'title': title, 'layer': layer, 'summary': summary,
                'source': src, 'format': 'converted from xlsx, content unchanged'}
        md = convert(os.path.join(BASE, src), meta, title)
        open(os.path.join(OUT, dst), 'w', encoding='utf-8').write(md)
        print(f'{dst:32s} {len(md):8d} 字符')
